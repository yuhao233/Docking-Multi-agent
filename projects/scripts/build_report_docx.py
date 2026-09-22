#!/usr/bin/env python3
"""把 ``docs/技术报告.md`` 渲染成专业排版的 Word（``docs/技术报告.docx``），并尽量转出 PDF。

设计要点
--------
* **只读 Markdown**：本脚本不修改 ``docs/技术报告.md``；插图采用「章节 → 图片」映射
  （见 :data:`FIGURE_MAP`），在渲染期把图片块插到指定小节的第 N 个正文块之后，
  保证每张图前面都有一段正文引出它。
* **支持的 Markdown 子集**：标题 1–4 级（``#``–``####``）、正文段落（软换行按 CJK 规则拼接）、
  无序/有序列表、GFM 表格、围栏代码块、行内 ``代码``、**加粗**、*斜体*、``>`` 引用块、
  ``![题注](路径)`` 图片。
* **排版**：封面页 + Word TOC 域（带缓存结果与真实页码）+ 页眉页脚（PAGE/NUMPAGES 域）+
  正文 12pt/1.5 倍行距/首行缩进 2 字符 + 中文 Noto Sans CJK SC + 等宽 Noto Sans Mono CJK SC。
* **两遍生成**：第一遍生成目录（页码留空）并转 PDF，从 PDF 里抓出每个标题的真实页码，
  第二遍把页码写进目录域的缓存结果 —— PDF 里目录是真页码，Word 里按 F9 仍可更新。
  两遍的目录条目完全一致，因此分页不变。
* **不溢出**：表格用固定布局 + 按内容权重分配列宽 + 按列数/内容长度选择字号；代码块用 PIL
  实测等宽字体宽度反推可用字号。
* **OOXML 元素顺序**：所有手工插入的 ``w:pPr``/``w:rPr``/``w:tcPr``/``w:trPr``/``w:tblPr``
  子元素都按 schema 顺序插入（Word 对顺序敏感，乱序会被判为文件损坏）。

用法::

    export UV_CACHE_DIR=/home/biolab/Tools/docking-agent/.uv-cache
    .venv/bin/python scripts/build_report_docx.py            # 生成 docx + pdf + 自检
    .venv/bin/python scripts/build_report_docx.py --no-pdf   # 只生成 docx
    .venv/bin/python scripts/build_report_docx.py --dry-run  # 只打印插图位置

依赖：``python-docx``（开发工具依赖，未写入运行时依赖）、``Pillow``、``soffice``/``pdfinfo``/
``pdftotext``（缺失时自动降级为「只出 docx」）。
"""

from __future__ import annotations

import argparse
import datetime
import re
import shutil
import subprocess
import sys
from typing import Any
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

try:  # pragma: no cover - 依赖缺失时给出可读报错
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:  # pragma: no cover
    sys.exit(f"需要 python-docx：uv pip install --python .venv/bin/python python-docx\n{exc}")

try:
    from PIL import Image, ImageFont
except ImportError:  # pragma: no cover
    Image = None
    ImageFont = None


REPO = Path(__file__).resolve().parent.parent

# --------------------------------------------------------------------------- #
# 常量
# --------------------------------------------------------------------------- #
BODY_EA = "Noto Sans CJK SC"          # 思源黑体等效（系统已装 Noto Sans CJK SC）
BODY_LATIN = "Noto Sans CJK SC"
MONO_FONT = "Noto Sans Mono CJK SC"
MONO_TTC = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
MONO_TTC_INDEX = 7                    # ttc 内 "Noto Sans Mono CJK SC" 的面序号

PAGE_W_CM, PAGE_H_CM = 21.0, 29.7     # A4
MARGIN_TB_CM, MARGIN_LR_CM = 2.2, 2.4
USABLE_W_CM = PAGE_W_CM - 2 * MARGIN_LR_CM          # 16.2 cm
USABLE_W_PT = USABLE_W_CM / 2.54 * 72               # ≈ 459.2 pt
FIG_W_CM = round(USABLE_W_CM * 0.85, 2)             # 图片统一宽度 = 可用宽度 85%
FIG_MAX_H_CM = 17.0

COLOR_TITLE = "1F3864"
COLOR_H2 = "1F3864"
COLOR_H3 = "2E5496"
COLOR_H4 = "333333"
COLOR_CAPTION = "404040"
CODE_BG = "F2F2F2"
QUOTE_BG = "EEF3FA"
QUOTE_BAR = "4472C4"
TBL_HEAD_BG = "DCE6F1"
TBL_ZEBRA_BG = "F7F9FC"
VERSION_FALLBACK = "0.6.0"

# ---- OOXML schema 子元素顺序（Word 对顺序敏感） ----
PPR_SEQ = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
           "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
           "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
           "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
           "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
           "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
           "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
           "w:sectPr", "w:pPrChange")
TCPR_SEQ = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
            "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
            "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
TRPR_SEQ = ("w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter",
            "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc",
            "w:hidden", "w:ins", "w:del", "w:trPrChange")
TBLPR_SEQ = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
             "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
             "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
             "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def insert_ordered(parent, elm, seq) -> None:
    """按 schema 顺序把 elm 插到 parent 里。"""
    name = "w:" + elm.tag.split("}")[1]
    if name in seq:
        successors = seq[seq.index(name) + 1:]
        parent.insert_element_before(elm, *successors)
    else:
        parent.append(elm)


# --------------------------------------------------------------------------- #
# Markdown 解析
# --------------------------------------------------------------------------- #
@dataclass
class Block:
    kind: str                      # heading|para|list|quote|code|table|image|consumed
    text: str = ""
    level: int = 0                 # 标题级别 / 列表缩进级别
    ordered: bool = False
    header: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    lines: list = field(default_factory=list)
    src_line: int = 0
    image: str = ""
    caption: str = ""


HEAD_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
IMAGE_RE = re.compile(r"^!\[(?P<cap>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
HR_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")


def is_table_sep(line: str) -> bool:
    s = line.strip()
    if not s.startswith("|"):
        return False
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", s)) and "-" in s


def split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_cjk(ch: str) -> bool:
    if not ch:
        return False
    o = ord(ch)
    if 0x3000 <= o <= 0x303F or 0x4E00 <= o <= 0x9FFF or 0xFF00 <= o <= 0xFFEF:
        return True
    if 0x3400 <= o <= 0x4DBF or 0x20000 <= o <= 0x2FA1F:
        return True
    return unicodedata.east_asian_width(ch) in ("W", "F")


def join_wrapped(lines: list[str]) -> str:
    """Markdown 软换行：中文之间不补空格，西文之间补一个空格。"""
    out = ""
    for raw in lines:
        piece = raw.strip()
        if not piece:
            continue
        if not out:
            out = piece
        elif _is_cjk(out[-1]) and _is_cjk(piece[0]):
            out += piece
        else:
            out += " " + piece
    return out


def parse_markdown(text: str) -> list[Block]:
    lines = text.split("\n")
    blocks: list[Block] = []
    i, n = 0, len(lines)
    para: list[str] = []
    para_line = 0

    def flush() -> None:
        nonlocal para, para_line
        if para:
            blocks.append(Block(kind="para", text=join_wrapped(para), src_line=para_line))
            para = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:                                     # 空行
            flush()
            i += 1
            continue

        if stripped.startswith("```"):                       # 围栏代码块
            flush()
            start = i + 1
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(Block(kind="code", lines=buf, src_line=start))
            continue

        if HR_RE.match(stripped):                            # 分隔线
            flush()
            blocks.append(Block(kind="hr", src_line=i + 1))
            i += 1
            continue

        m = HEAD_RE.match(line)                              # 标题
        if m:
            flush()
            blocks.append(Block(kind="heading", text=m.group(2).strip(),
                                level=len(m.group(1)), src_line=i + 1))
            i += 1
            continue

        if stripped.startswith(">"):                         # 引用块
            flush()
            buf = []
            start = i
            while i < n and lines[i].strip().startswith(">"):
                buf.append(QUOTE_RE.match(lines[i]).group(1))
                i += 1
            blocks.append(Block(kind="quote", text=join_wrapped(buf), src_line=start + 1))
            continue

        if stripped.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):   # 表格
            flush()
            start = i
            header = split_row(lines[i])
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                while len(cells) < len(header):
                    cells.append("")
                rows.append(cells[: len(header)])
                i += 1
            blocks.append(Block(kind="table", header=header, rows=rows, src_line=start + 1))
            continue

        m = BULLET_RE.match(line)                            # 无序列表
        if m:
            flush()
            while i < n:
                b = BULLET_RE.match(lines[i])
                if not b:
                    break
                blocks.append(Block(kind="list", text=b.group(2).strip(),
                                    level=min(len(b.group(1)) // 2, 2), ordered=False,
                                    src_line=i + 1))
                i += 1
            continue

        m = ORDERED_RE.match(line)                           # 有序列表
        if m:
            flush()
            while i < n:
                o = ORDERED_RE.match(lines[i])
                if not o:
                    break
                blocks.append(Block(kind="list", text=f"{o.group(2)}. {o.group(3).strip()}",
                                    level=min(len(o.group(1)) // 2, 2), ordered=True,
                                    src_line=i + 1))
                i += 1
            continue

        m = IMAGE_RE.match(stripped)                         # 图片
        if m:
            flush()
            blocks.append(Block(kind="image", image=m.group("src"), caption=m.group("cap"),
                                src_line=i + 1))
            i += 1
            continue

        if not para:
            para_line = i + 1
        para.append(line)
        i += 1

    flush()
    return blocks


# --------------------------------------------------------------------------- #
# 行内解析
# --------------------------------------------------------------------------- #
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*|~~.+?~~)")


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    strike: bool = False


def parse_inline(text: str) -> list[Run]:
    """行内标记 → run 列表。

    非 http(s) 的 ``[文字](目标)``（例如 ``[Na+](反离子)``）按普通文本原样保留：
    报告正文不允许出现 URL，也不应把这类写法误当链接。
    """
    text = LINK_RE.sub(lambda m: m.group(1) if re.match(r"^[a-z]+://", m.group(2)) else m.group(0),
                       text)
    runs: list[Run] = []
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            runs.append(Run(token[2:-2], bold=True))
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            runs.append(Run(token[1:-1], code=True))
        elif token.startswith("~~") and token.endswith("~~") and len(token) > 4:
            runs.append(Run(token[2:-2], strike=True))
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            runs.append(Run(token[1:-1], italic=True))
        else:
            runs.append(Run(token))
    return runs


def plain_text(text: str) -> str:
    """去掉 Markdown 标记的纯文本（用于目录/题注）。"""
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    t = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", t)
    t = t.replace("`", "")
    return t.strip()


# --------------------------------------------------------------------------- #
# 插图映射（章节 → 图片）
# --------------------------------------------------------------------------- #
# 配图**入库**在 `docs/report-media/`（审计 3.5）：此前 13/13 张都指向被 gitignore 的
# `var/` 运行产物，干净克隆上重跑生成器只会得到 13 个红色「［缺图］」占位符，
# 而提交的 DOCX 里却有真图 —— 属于「产物无法从仓库复现」。
# 图片是界面截图与报告图表的**副本**，与运行目录解耦；重新截图后用
# `scripts/ui_shot.py` / `scripts/browser_check.py` 覆盖同名文件即可。
# 缺图时生成器**默认拒绝出文档**（见 main 的 --allow-missing-figures）。


@dataclass
class Figure:
    anchor: str          # 触发标题的正则（匹配标题正文）
    after: int           # 该小节内第 N 个正文块之后插入
    path: str
    caption: str
    used: bool = False


FIGURE_MAP: list[Figure] = [
    Figure(r"^2\.\s*总体架构", 1, "docs/report-media/ui-workbench-chat.png",
           "工作台全貌（对话模式首屏：左「任务配置」、中「多 Agent 编排」实时编排图、下「执行区」）"),
    Figure(r"^3\.1\s", 1, "docs/report-media/ui-agent-running.png",
           "一次真实运行中的编排时间轴、实时分子结果与工具调用轨迹（SSE 推送渲染）"),
    Figure(r"^4\.5\s", 1, "docs/report-media/ui-chat-attached.png",
           "对话模式下的附件上传与 @ 引用（点选附件后以同一 conversation_id 续问）"),
    Figure(r"^7\.2\s", 1, "docs/report-media/ui-manual-params.png",
           "参数模式：受体与对接盒（center / size）的显式指定；留空则走工具定盒"),
    Figure(r"^7\.3\s", 1, "docs/report-media/ui-more-params.png",
           "「更多参数」面板：结合位点中心 X / Y / Z 的显式输入（留空即由口袋工具确定）"),
    Figure(r"^8\.3\s", 1, "docs/report-media/ui-exhaustiveness-preset.png",
           "搜索强度一键预设（快速初筛 / 平衡 / 高精度 / 恢复系统默认）与自动规划出的 EXHAUSTIVENESS: 24"),
    Figure(r"^9\.\s*评估与报告", 1, "docs/report-media/chart-docking.png",
           "对接亲和力对比图（含阳性对照；来自实测运行 20260916-085944-4688 的报告产物）"),
    Figure(r"^9\.1\s", 1, "docs/report-media/chart-property.png",
           "理化性质空间（横轴分子量、纵轴 logP、颜色/点大小 = 对接强度，虚线为 Lipinski 边界）"),
    Figure(r"^9\.4\s", 1, "docs/report-media/ui-pdf-cover.png",
           "系统自动生成的 PDF 报告第 1 页：封面、工具版本、任务与参数（真实运行产物）"),
    Figure(r"^9\.5\s", 1, "docs/report-media/chart-structure-grid.png",
           "候选分子结构网格（按对接亲和力排序，左一为阳性对照）"),
    Figure(r"^9\.5\s", 2, "docs/report-media/chart-affinity-histogram.png",
           "亲和力分布直方图（绿色 / 紫色 / 红色虚线分别为均值、中位数与阳性对照）"),
    Figure(r"^12\.3\s", 1, "docs/report-media/ui-settings-crop.png",
           "设置页：模型与部署级参数（局部视图；含 API 地址与密钥的字段已裁去）"),
    Figure(r"^12\.5\s", 1, "docs/report-media/ui-run-detail.png",
           "运行详情页：阶段日志、结果总览与产物下载（报告 PDF / MD / CSV / 位姿 ZIP）"),
]


def missing_figure_files(repo: Path, figures: list[Figure] | None = None) -> list[str]:
    """返回 `FIGURE_MAP` 中在仓库里**不存在**的配图路径。

    这是「干净克隆能否复现 DOCX/PDF」的判据：配图必须入库（`docs/report-media/`），
    不能指向被 gitignore 的 `var/` 运行产物。缺图时 `main()` 默认拒绝出文档。
    """
    return [fig.path for fig in (figures if figures is not None else FIGURE_MAP)
            if not (repo / fig.path).is_file()]


def insert_figures(blocks: list[Block], figures: list[Figure], dry_run: bool = False) -> list[Block]:
    """把小节第 N 个正文块之后的位置插进图片块（每图前必有一段正文）。"""
    out: list[Block] = []
    heading: str | None = None
    seen = 0
    for blk in blocks:
        out.append(blk)
        if blk.kind == "heading":
            heading = blk.text
            seen = 0
            continue
        seen += 1
        if heading is None:
            continue
        for fig in figures:
            if fig.used or fig.after != seen or not re.search(fig.anchor, heading):
                continue
            if dry_run:
                print(f"  [fig] 「{heading}」第 {seen} 块后 → {fig.path}")
            out.append(Block(kind="image", image=fig.path, caption=fig.caption))
            fig.used = True
    return out


BOLD_LEADIN_RE = re.compile(r"^\*\*.+?\*\*.*[：:]\s*$")


def attach_table_captions(blocks: list[Block]) -> None:
    """给每个表格配编号题注。

    紧邻表格之前的「短加粗引导句」（如 ``**`AUTO_PARAM_*` 全部默认值总表**：``）会被**改写成**
    该表的题注（``表 N 题目``），既不重复也不丢文字；否则用最近的小节标题（去掉「N.N」编号）
    作为题目，文档最前面的无标题表格用「文档基本信息」。
    """
    heading = "文档基本信息"
    for idx, blk in enumerate(blocks):
        if blk.kind == "heading":
            if blk.level >= 2:                       # md '#' 是文档大标题，不作表格题注
                heading = re.sub(r"^\d+(\.\d+)*\s*", "", plain_text(blk.text)).strip()
            continue
        if blk.kind != "table":
            continue
        title = None
        prev = blocks[idx - 1] if idx else None
        if prev is not None and prev.kind == "para" and BOLD_LEADIN_RE.match(prev.text.strip()):
            plain = plain_text(prev.text).rstrip("：: ").strip()
            if 0 < len(plain) <= 80:
                title = plain
                prev.kind = "consumed"
        blk.caption = title or heading or "表格"


# --------------------------------------------------------------------------- #
# 样式与底层工具
# --------------------------------------------------------------------------- #
def set_style_font(style, size=None, bold=None, color=None, italic=None,
                   latin=BODY_LATIN, ea=BODY_EA) -> None:
    style.font.name = latin
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), ea)
    rfonts.set(qn("w:cs"), latin)


def set_run_font(run, size=None, bold=None, color=None, italic=None, code=False) -> None:
    latin = MONO_FONT if code else BODY_LATIN
    ea = MONO_FONT if code else BODY_EA
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), ea)
    rfonts.set(qn("w:cs"), latin)


def style_or_add(doc, name: str, base: str = "Normal") -> Any:
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
        return st


def para_indent_zero(p) -> None:
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)


def first_line_indent_2chars(p) -> None:
    """首行缩进 2 字符：同时写 firstLineChars（Word 语义）与 firstLine（LibreOffice 语义）。"""
    p.paragraph_format.first_line_indent = Pt(24)
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), "200")
        ind.set(qn("w:firstLine"), "480")     # 24pt = 2 × 12pt


def para_shade(p, fill: str) -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, PPR_SEQ)


def para_border(p, edges=("top", "left", "bottom", "right"), sz=6, color="D0D7E5",
                space=6) -> None:
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        insert_ordered(pPr, bdr, PPR_SEQ)
    for e in edges:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        bdr.append(el)


def add_field(paragraph, instr: str, placeholder: str = "1", size=None, color=None) -> None:
    """插入一个简单域（PAGE / NUMPAGES），并带一个缓存结果。"""
    def _fld(kind):
        r = paragraph.add_run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r._r.append(fc)
        set_run_font(r, size=size, color=color)
        return r
    _fld("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    set_run_font(r, size=size, color=color)
    _fld("separate")
    r = paragraph.add_run(placeholder)
    set_run_font(r, size=size, color=color)
    _fld("end")


def tbl_borders(table, color="9DB2CE", sz=6) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_SEQ)


def tbl_fixed_layout(table) -> None:
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    insert_ordered(tblPr, layout, TBLPR_SEQ)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 40), ("left", 85), ("bottom", 40), ("right", 85)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    insert_ordered(tblPr, mar, TBLPR_SEQ)


def tbl_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    insert_ordered(trPr, cant, TRPR_SEQ)
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    insert_ordered(trPr, hdr, TRPR_SEQ)


def cell_shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd, TCPR_SEQ)


# --------------------------------------------------------------------------- #
# 渲染器
# --------------------------------------------------------------------------- #
class ReportBuilder:
    def __init__(self, md_path: Path, out_path: Path, page_map: dict[str, int] | None = None,
                 toc_field: bool = True) -> None:
        self.md_path = md_path
        self.out_path = out_path
        self.page_map = page_map or {}
        self.toc_field = toc_field
        self.repo = md_path.parent.parent
        self.doc = Document()
        self.fig_no = 0
        self.tbl_no = 0
        self.headings: list[tuple[int, str]] = []
        self.missing_images: list[str] = []
        self.used_images: list[tuple[str, str]] = []
        self._setup()

    # ---------------- 页面与样式 ----------------
    def _setup(self) -> None:
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
        sec.top_margin = sec.bottom_margin = Cm(MARGIN_TB_CM)
        sec.left_margin = sec.right_margin = Cm(MARGIN_LR_CM)
        sec.header_distance = Cm(1.2)
        sec.footer_distance = Cm(1.1)
        sec.different_first_page_header_footer = True     # 封面不显示页眉页脚

        st = doc.styles["Normal"]
        set_style_font(st, size=12)
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.space_after = Pt(4)
        pf.space_before = Pt(0)

        for name, size, bold, color, sb, sa in (
            ("Heading 1", 22, True, COLOR_TITLE, 24, 12),
            ("Heading 2", 17, True, COLOR_H2, 18, 10),
            ("Heading 3", 14, True, COLOR_H3, 14, 8),
            ("Heading 4", 12.5, True, COLOR_H4, 12, 6),
            ("Heading 5", 12, True, COLOR_H4, 10, 6),
            ("Heading 6", 12, True, COLOR_H4, 8, 4),
        ):
            s = doc.styles[name]
            set_style_font(s, size=size, bold=bold, color=color)
            p = s.paragraph_format
            p.first_line_indent = Pt(0)
            p.left_indent = Pt(0)
            p.space_before = Pt(sb)
            p.space_after = Pt(sa)
            p.line_spacing = 1.25
            p.keep_with_next = True

        for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
            s = style_or_add(doc, name, base="Normal")
            set_style_font(s, size=12)
            p = s.paragraph_format
            p.line_spacing = 1.4
            p.space_after = Pt(3)

        s = style_or_add(doc, "ReportCaption")
        set_style_font(s, size=9.5, color=COLOR_CAPTION)
        p = s.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.first_line_indent = Pt(0)
        p.line_spacing = 1.15
        p.space_before = Pt(4)
        p.space_after = Pt(12)

        s = style_or_add(doc, "ReportTableCaption")
        set_style_font(s, size=10, bold=True, color=COLOR_H2)
        p = s.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.first_line_indent = Pt(0)
        p.line_spacing = 1.15
        p.space_before = Pt(10)
        p.space_after = Pt(3)
        p.keep_with_next = True

        s = style_or_add(doc, "ReportCode")
        set_style_font(s, size=8.5, latin=MONO_FONT, ea=MONO_FONT)
        p = s.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.first_line_indent = Pt(0)
        p.left_indent = Cm(0.2)
        p.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.space_before = Pt(6)
        p.space_after = Pt(10)

        s = style_or_add(doc, "ReportQuote")
        set_style_font(s, size=10.5, color="333333")
        p = s.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.first_line_indent = Pt(0)
        p.left_indent = Cm(0.5)
        p.right_indent = Cm(0.2)
        p.line_spacing = 1.3
        p.space_before = Pt(6)
        p.space_after = Pt(10)

        for name, size, indent in (("ReportTOC1", 11, 0.0), ("ReportTOC2", 10.5, 0.7)):
            s = style_or_add(doc, name)
            set_style_font(s, size=size, bold=(name == "ReportTOC1"))
            p = s.paragraph_format
            p.first_line_indent = Pt(0)
            p.left_indent = Cm(indent)
            p.line_spacing = 1.2
            p.space_before = Pt(2 if name == "ReportTOC1" else 0)
            p.space_after = Pt(0)
            p.tab_stops.add_tab_stop(Cm(USABLE_W_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        self._header_footer(sec)

    def _header_footer(self, sec) -> None:
        p = sec.header.paragraphs[0]
        p.text = ""
        para_indent_zero(p)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(USABLE_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        set_run_font(p.add_run("多 Agent 协同小分子对接筛选系统 · 技术报告"), size=8.5, color="808080")
        set_run_font(p.add_run("\t"), size=8.5, color="808080")
        set_run_font(p.add_run(f"v{self.version()}"), size=8.5, color="808080")
        para_border(p, edges=("bottom",), sz=4, color="BFBFBF", space=2)

        p = sec.footer.paragraphs[0]
        p.text = ""
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        set_run_font(p.add_run("第 "), size=8.5, color="808080")
        add_field(p, "PAGE", "1", size=8.5, color="808080")
        set_run_font(p.add_run(" 页 / 共 "), size=8.5, color="808080")
        add_field(p, "NUMPAGES", "1", size=8.5, color="808080")
        set_run_font(p.add_run(" 页"), size=8.5, color="808080")

    def version(self) -> str:
        init = self.repo / "src" / "docking_agent" / "__init__.py"
        try:
            m = re.search(r'__version__\s*=\s*"([^"]+)"', init.read_text(encoding="utf-8"))
            if m:
                return m.group(1)
        except OSError as _exc:  # noqa: BLE001
            print(f"[build_report_docx] 忽略可恢复错误：{_exc}", file=sys.stderr)
        return VERSION_FALLBACK

    # ---------------- 封面 ----------------
    def _cover(self, summary: str) -> None:
        doc = self.doc
        for _ in range(4):
            p = doc.add_paragraph()
            para_indent_zero(p)
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(""), size=12)

        p = doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run("多 Agent 协同小分子对接筛选系统"), size=28, bold=True, color=COLOR_TITLE)

        p = doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run_font(p.add_run("技 术 报 告"), size=22, bold=True, color=COLOR_TITLE)

        p = doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        para_border(p, edges=("bottom",), sz=12, color=COLOR_H3, space=1)

        for text, size, color, after in (
            ("本地部署版 · 自动准备 / 对接 / 评估 / 异常处理 / 报告", 15, COLOR_H3, 10),
            (f"版本 v{self.version()}", 12, "404040", 2),
            (f"生成日期 {datetime.date.today().isoformat()}", 12, "404040", 20),
        ):
            p = doc.add_paragraph()
            para_indent_zero(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(after)
            set_run_font(p.add_run(text), size=size, color=color)

        p = doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.3)
        p.paragraph_format.right_indent = Cm(1.3)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        para_shade(p, QUOTE_BG)
        para_border(p, sz=4, color="C7D5EA", space=8)
        set_run_font(p.add_run("摘要　"), size=10.5, bold=True, color=COLOR_H2)
        set_run_font(p.add_run(summary), size=10.5, color="333333")

        for _ in range(4):
            p = doc.add_paragraph()
            para_indent_zero(p)
            p.paragraph_format.space_after = Pt(0)

        for text in (
            "内容源：docs/技术报告.md",
            "排版与插图：scripts/build_report_docx.py 自动生成",
            "本报告全部数值来自仓库内代码、既有文档与实测运行；未实测项已明确标注。",
        ):
            p = doc.add_paragraph()
            para_indent_zero(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(text), size=9, color="707070")

        self._page_break()

    def _page_break(self) -> None:
        p = self.doc.add_paragraph()
        para_indent_zero(p)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run._r.append(br)

    # ---------------- 目录 ----------------
    def _toc(self, headings: list[tuple[int, str]]) -> None:
        p = self.doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run_font(p.add_run("目　　录"), size=20, bold=True, color=COLOR_TITLE)

        styles = {1: "ReportTOC1", 2: "ReportTOC2"}
        entries = list(headings)
        for idx, (lvl, txt) in enumerate(entries):
            size = 11 if lvl == 1 else 10.5
            ep = self.doc.add_paragraph(style=styles.get(lvl, "ReportTOC2"))
            para_indent_zero(ep)
            ep.paragraph_format.left_indent = Cm(0.0 if lvl == 1 else 0.7)
            ep.paragraph_format.tab_stops.add_tab_stop(
                Cm(USABLE_W_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if self.toc_field and idx == 0:
                r = ep.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "begin")
                r._r.append(fc)
                r = ep.add_run()
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-3" \\h \\z \\u '
                r._r.append(it)
                r = ep.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "separate")
                r._r.append(fc)
            set_run_font(ep.add_run(txt), size=size, bold=(lvl == 1))
            set_run_font(ep.add_run("\t" + str(self.page_map.get(txt, ""))), size=size, bold=(lvl == 1))
            if self.toc_field and idx == len(entries) - 1:
                r = ep.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "end")
                r._r.append(fc)

        note = self.doc.add_paragraph()
        para_indent_zero(note)
        note.paragraph_format.space_before = Pt(14)
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(note.add_run(
            "提示：本目录为 Word 目录域（TOC）的缓存结果；在 Word 中按 Ctrl+A → F9 或右键「更新域」可刷新页码。"),
            size=8.5, color="808080")
        self._page_break()

    # ---------------- 正文 ----------------
    def render(self, blocks: list[Block], summary: str) -> dict:
        cp = self.doc.core_properties
        cp.title = "多 Agent 协同小分子对接筛选系统 · 技术报告"
        cp.author = "docking-agent 项目组"
        cp.subject = "本地部署版：自动准备 / 对接 / 评估 / 异常处理 / 报告"
        cp.comments = "由 scripts/build_report_docx.py 从 docs/技术报告.md 自动排版生成"
        cp.category = "技术报告"
        self._cover(summary)
        body = [b for b in blocks if not (b.kind == "heading" and b.level == 1)]
        self.headings = self._collect_headings(body)
        self._toc(self.headings)
        for blk in body:
            self._block(blk)
        self.doc.save(str(self.out_path))
        return self._stats()

    def _collect_headings(self, blocks: list[Block]) -> list[tuple[int, str]]:
        out = []
        for b in blocks:
            if b.kind == "heading":
                lvl = b.level - 1            # md '##' → 目录第 1 级
                if lvl in (1, 2):
                    out.append((lvl, plain_text(b.text)))
        return out

    def _block(self, blk: Block) -> None:
        if blk.kind == "heading":
            self._heading(blk)
        elif blk.kind == "para":
            self._paragraph(blk.text)
        elif blk.kind == "list":
            self._list(blk)
        elif blk.kind == "quote":
            self._quote(blk.text)
        elif blk.kind == "code":
            self._code(blk.lines)
        elif blk.kind == "table":
            self._table(blk)
        elif blk.kind == "image":
            self._figure(blk.image, blk.caption)
        elif blk.kind == "hr":
            p = self.doc.add_paragraph()
            para_indent_zero(p)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            set_run_font(p.add_run(""), size=2)
            para_border(p, edges=("bottom",), sz=4, color="C7D5EA", space=1)
        # 'consumed'：已改写成表格题注，跳过

    def _heading(self, blk: Block) -> None:
        level = min(blk.level, 6)
        p = self.doc.add_paragraph(style=f"Heading {level}")
        para_indent_zero(p)
        if blk.level == 2:                     # 一级小节（## N. xxx）前分页
            p.paragraph_format.page_break_before = True
        size = {1: 22, 2: 17, 3: 14, 4: 12.5, 5: 12, 6: 12}[level]
        color = {1: COLOR_TITLE, 2: COLOR_H2, 3: COLOR_H3}.get(level, COLOR_H4)
        self._add_inline(p, blk.text, base_size=size, base_bold=True, base_color=color)

    def _paragraph(self, text: str, size: float | None = None) -> None:
        p = self.doc.add_paragraph()
        para_indent_zero(p)
        first_line_indent_2chars(p)
        self._add_inline(p, text, base_size=size)

    def _list(self, blk: Block) -> None:
        style = "List Bullet"
        if blk.level >= 1:
            style = f"List Bullet {min(blk.level + 1, 3)}"
        p = self.doc.add_paragraph(style=style)
        self._add_inline(p, blk.text)

    def _quote(self, text: str) -> None:
        p = self.doc.add_paragraph(style="ReportQuote")
        para_indent_zero(p)
        p.paragraph_format.left_indent = Cm(0.5)
        para_shade(p, QUOTE_BG)
        para_border(p, edges=("left",), sz=18, color=QUOTE_BAR, space=8)
        self._add_inline(p, text, base_size=10.5)

    def _code(self, lines: list[str]) -> None:
        size = self._code_font_size(lines)
        p = self.doc.add_paragraph(style="ReportCode")
        para_indent_zero(p)
        p.paragraph_format.left_indent = Cm(0.2)
        para_shade(p, CODE_BG)
        para_border(p, sz=4, color="DDDDDD", space=6)
        for i, line in enumerate(lines):
            if i:
                br = p.add_run()
                br._r.append(OxmlElement("w:br"))
            r = p.add_run(line if line else " ")
            set_run_font(r, size=size, code=True, color="1A1A1A")

    def _code_font_size(self, lines: list[str]) -> float:
        if not lines:
            return 8.5
        avail = USABLE_W_PT - 26          # 减去缩进、边框与内边距
        maxw = max((self._em_width(l) for l in lines), default=1.0)
        size = avail / maxw if maxw else 8.5
        return max(5.5, min(8.5, round(size, 2)))

    _FONT_CACHE: dict = {}

    @classmethod
    def _mono_font(cls):
        if "f" not in cls._FONT_CACHE:
            f = None
            if ImageFont is not None and Path(MONO_TTC).exists():
                for idx in (MONO_TTC_INDEX, 0):
                    try:
                        f = ImageFont.truetype(MONO_TTC, 200, index=idx)
                        break
                    except Exception:
                        f = None
            cls._FONT_CACHE["f"] = f
        return cls._FONT_CACHE["f"]

    @classmethod
    def _em_width(cls, text: str) -> float:
        """文本按 1em 计的宽度（等宽字体 PIL 实测；无 PIL 时按字符宽度估算）。"""
        f = cls._mono_font()
        if f is not None:
            try:
                return max(f.getlength(text) / 200.0, 0.6)
            except Exception as _exc:  # noqa: BLE001
                print(f"[build_report_docx] 忽略可恢复错误：{_exc}", file=sys.stderr)
        return sum(2.0 if unicodedata.east_asian_width(c) in ("W", "F", "A") else 0.6 for c in text)

    def _table(self, blk: Block) -> None:
        self.tbl_no += 1
        ncols = len(blk.header)
        size = self._table_font_size(blk, ncols)

        cap = self.doc.add_paragraph(style="ReportTableCaption")
        para_indent_zero(cap)
        set_run_font(cap.add_run(f"表 {self.tbl_no}　"), size=10, bold=True, color=COLOR_H2)
        self._add_inline(cap, blk.caption or "表格", base_size=10, base_bold=True,
                         base_color=COLOR_H2)

        table = self.doc.add_table(rows=1, cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_fixed_layout(table)
        tbl_borders(table)

        widths = self._column_widths(blk, ncols)
        hdr = table.rows[0]
        for j, cell_text in enumerate(blk.header):
            cell = hdr.cells[j]
            cell.width = Cm(widths[j])
            cell_shade(cell, TBL_HEAD_BG)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            para_indent_zero(p)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(0)
            self._add_inline(p, cell_text, base_size=size, base_bold=True, base_color=COLOR_H2)
        tbl_repeat_header(hdr)

        for i, row in enumerate(blk.rows):
            tr = table.add_row()
            for j, cell_text in enumerate(row[:ncols]):
                cell = tr.cells[j]
                cell.width = Cm(widths[j])
                if i % 2 == 1:
                    cell_shade(cell, TBL_ZEBRA_BG)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                p = cell.paragraphs[0]
                para_indent_zero(p)
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(0)
                self._add_inline(p, cell_text, base_size=size)

        tail = self.doc.add_paragraph()
        para_indent_zero(tail)
        tail.paragraph_format.space_after = Pt(2)
        tail.paragraph_format.line_spacing = 1.0

    def _table_font_size(self, blk: Block, ncols: int) -> float:
        longest = 0
        for cell in blk.header + [c for row in blk.rows for c in row]:
            longest = max(longest, len(plain_text(cell)))
        if ncols <= 3:
            size = 10.5
        elif ncols == 4:
            size = 9.5
        elif ncols == 5:
            size = 8.5
        else:
            size = 8.5
        if longest > 160 and size > 8.5:
            size -= 0.5
        return max(8.0, size)

    def _column_widths(self, blk: Block, ncols: int) -> list[float]:
        """按内容量 + 最长不可断词给列分配宽度（总宽固定 = 可用宽度，绝不溢出）。"""
        total = USABLE_W_CM
        weights = []
        for j in range(ncols):
            vals = [plain_text(blk.header[j]) if j < len(blk.header) else ""]
            vals += [plain_text(row[j]) if j < len(row) else "" for row in blk.rows]
            mx = max((len(v) for v in vals), default=1)
            avg = sum(len(v) for v in vals) / max(len(vals), 1)
            longest_token = 0
            for v in vals:
                for tok in re.split(r"[\s、，,；;（）()/｜|]+", v):
                    longest_token = max(longest_token, len(tok))
            weight = 0.35 * mx + 0.45 * avg + 0.35 * longest_token * 3.0
            weights.append(max(3.0, weight ** 0.8))
        s = sum(weights)
        widths = [max(1.35, total * w / s) for w in weights]
        cap = total * 0.62
        widths = [min(w, cap) for w in widths]
        s2 = sum(widths)
        return [round(w * total / s2, 2) for w in widths]

    def _figure(self, rel: str, caption: str) -> None:
        path = Path(rel) if Path(rel).is_absolute() else (self.repo / rel)
        if not path.exists():
            self.missing_images.append(rel)
            p = self.doc.add_paragraph()
            para_indent_zero(p)
            set_run_font(p.add_run(f"［缺图：{rel}］"), size=10, color="C00000")
            return
        self.fig_no += 1
        width_cm = FIG_W_CM
        if Image is not None:
            try:
                with Image.open(path) as im:
                    w, h = im.size
                if h and w:
                    height_cm = width_cm * h / w
                    if height_cm > FIG_MAX_H_CM:
                        width_cm = FIG_MAX_H_CM * w / h
            except Exception as _exc:  # noqa: BLE001
                print(f"[build_report_docx] 忽略可恢复错误：{_exc}", file=sys.stderr)
        p = self.doc.add_paragraph()
        para_indent_zero(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Cm(width_cm))

        cap = self.doc.add_paragraph(style="ReportCaption")
        para_indent_zero(cap)
        cap.paragraph_format.space_after = Pt(14)
        set_run_font(cap.add_run(f"图 {self.fig_no}　"), size=9.5, bold=True, color=COLOR_H2)
        self._add_inline(cap, caption, base_size=9.5, base_color=COLOR_CAPTION)
        self.used_images.append((rel, caption))

    def _add_inline(self, p, text: str, base_size=None, base_bold=False, base_color=None) -> None:
        for run in parse_inline(text):
            r = p.add_run(run.text)
            size = base_size
            if run.code and base_size:
                size = round(base_size * 0.92, 2)
            set_run_font(r, size=size, bold=base_bold or run.bold,
                         italic=True if run.italic else None,
                         color="8B1A1A" if run.code else base_color,
                         code=run.code)
            if run.strike:
                r.font.strike = True

    # ---------------- 自检 ----------------
    def _stats(self) -> dict:
        doc = Document(str(self.out_path))
        texts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    texts.append(c.text)
        joined = "\n".join(texts)
        return {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "inline_shapes": len(doc.inline_shapes),
            "bytes": self.out_path.stat().st_size,
            "urls": re.findall(r"https?://\S+", joined),
            "chars": len(joined),
        }


# --------------------------------------------------------------------------- #
# 媒体预处理
# --------------------------------------------------------------------------- #
def prepare_media(repo: Path) -> None:
    """生成派生图片：设置页截图裁掉含 API 地址 / 密钥的字段区域（正文不允许出现 URL）。"""
    outdir = repo / "var" / "tmp" / "report_media"
    outdir.mkdir(parents=True, exist_ok=True)
    src = repo / "var" / "tmp" / "ui_baseline" / "settings.png"
    dst = outdir / "settings_crop.png"
    if dst.exists() or not src.exists() or Image is None:
        return
    with Image.open(src) as im:
        w, h = im.size
        im.crop((int(w * 0.135), int(h * 0.30), w, h)).save(dst)


# --------------------------------------------------------------------------- #
# PDF 转换与页码提取
# --------------------------------------------------------------------------- #
def soffice_pdf(docx: Path, outdir: Path, profile: Path) -> Path | None:
    soffice = shutil.which("soffice") or "/usr/bin/soffice"
    if not Path(soffice).exists():
        print("  [warn] 未找到 soffice，跳过 PDF 转换")
        return None
    profile.mkdir(parents=True, exist_ok=True)
    cmd = [soffice, "--headless", "--norestore", "--invisible", "--nolockcheck",
           f"-env:UserInstallation=file://{profile}",
           "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(outdir), str(docx)]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=1200)
    except Exception as exc:
        print(f"  [warn] soffice 转换失败：{exc}")
        return None
    pdf = outdir / (docx.stem + ".pdf")
    return pdf if pdf.exists() else None


def pdf_pages(pdf: Path) -> int:
    try:
        out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, timeout=60).stdout
        m = re.search(r"^Pages:\s+(\d+)", out, re.M)
        return int(m.group(1)) if m else 0
    except Exception:
        return 0


def pdf_page_texts(pdf: Path, pages: int) -> list[str]:
    texts = []
    for i in range(1, pages + 1):
        try:
            out = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdf), "-"],
                                 capture_output=True, text=True, timeout=180).stdout
        except Exception:
            out = ""
        texts.append(re.sub(r"\s+", "", out))
    return texts


def build_page_map(pdf: Path, headings: list[tuple[int, str]]) -> dict[str, int]:
    """从 PDF 里抓每个标题所在页码（目录页由「更新域」提示定位）。"""
    pages = pdf_pages(pdf)
    if not pages:
        return {}
    texts = pdf_page_texts(pdf, pages)
    toc_last = 0
    for i, t in enumerate(texts[:10], start=1):
        if "更新域" in t:
            toc_last = i
    cur = toc_last + 1 if toc_last else 2
    mapping: dict[str, int] = {}
    for _, txt in headings:
        needle = re.sub(r"\s+", "", txt)
        found = None
        for i in range(cur - 1, pages):
            if needle and needle in texts[i]:
                found = i + 1
                break
        if found:
            mapping[txt] = found
            cur = found
        else:
            mapping[txt] = cur
    return mapping


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def extract_summary(md_text: str) -> str:
    m = re.search(r"\*\*一句话摘要\*\*[：:]\s*(.+?)(?:\n\s*\n)", md_text, re.S)
    if m:
        return plain_text(re.sub(r"\s*\n\s*", "", m.group(1)))
    return ("本系统在单机上把「小分子库 × 蛋白质受体」的真实分子对接筛选完整自动化："
            "主管 Agent 统筹 4 个专业子 Agent，用 RDKit / meeko / AutoDock Vina 等真实计算引擎"
            "产出可追溯的性质、对接能量、结合模式比较与固定 9 章报告。")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="把技术报告 Markdown 渲染为专业排版的 Word/PDF")
    ap.add_argument("--md", default=str(REPO / "docs" / "技术报告.md"))
    ap.add_argument("--out", default=str(REPO / "docs" / "技术报告.docx"))
    ap.add_argument("--no-pdf", action="store_true", help="只生成 docx")
    ap.add_argument("--plain-toc", action="store_true", help="目录不用 TOC 域，改为静态目录")
    ap.add_argument("--dry-run", action="store_true", help="只打印插图位置，不生成文件")
    ap.add_argument("--allow-missing-figures", action="store_true",
                    help="配图缺失时仍出文档（会在正文里留下红色「［缺图］」占位符）")
    args = ap.parse_args(argv)

    md_path = Path(args.md)
    out_path = Path(args.out)
    md_text = md_path.read_text(encoding="utf-8")
    summary = extract_summary(md_text)

    # 配图必须入库（docs/report-media/）：否则干净克隆上重跑只会得到「［缺图］」占位符，
    # 而提交的 DOCX 里却有真图 —— 产物无法从仓库复现（审计 3.5）。
    missing_files = missing_figure_files(REPO)
    if missing_files and not args.allow_missing_figures:
        print(f"❌ 配图缺失（{len(missing_files)}/{len(FIGURE_MAP)}）：{missing_files}", file=sys.stderr)
        print("   请把配图提交到 docs/report-media/，或修 FIGURE_MAP 的路径；"
              "确需出「带占位符」的文档时加 --allow-missing-figures。", file=sys.stderr)
        return 2

    blocks = parse_markdown(md_text)
    attach_table_captions(blocks)
    blocks = insert_figures(blocks, FIGURE_MAP, dry_run=args.dry_run)
    if args.dry_run:
        return 0

    prepare_media(REPO)
    work = REPO / "var" / "tmp" / "report_media"
    work.mkdir(parents=True, exist_ok=True)
    pass1 = work / "_pass1.docx"

    # ---- 第一遍：目录页码留空 → PDF → 抓真实页码 ----
    page_map: dict[str, int] = {}
    if not args.no_pdf:
        b1 = ReportBuilder(md_path, pass1, page_map={}, toc_field=not args.plain_toc)
        b1.render(_blocks(md_text), summary)
        pdf1 = soffice_pdf(pass1, work, work / "lo_profile")
        if pdf1:
            page_map = build_page_map(pdf1, b1.headings)
            print(f"  [toc] 从 {pdf1.name} 提取到 {len(page_map)}/{len(b1.headings)} 个标题页码")
        else:
            print("  [toc] 未生成 PDF，目录页码留空（Word 中按 F9 可更新）")

    # ---- 第二遍：写入页码，生成正式文件 ----
    builder = ReportBuilder(md_path, out_path, page_map=page_map, toc_field=not args.plain_toc)
    stats = builder.render(_blocks(md_text), summary)

    print(f"✅ 生成 {out_path.relative_to(REPO)}")
    print(f"   段落 {stats['paragraphs']} ｜ 表格 {stats['tables']} ｜ 内嵌图片 {stats['inline_shapes']} "
          f"｜ {stats['bytes'] / 1024:.1f} KB ｜ 文本 {stats['chars']} 字")
    print(f"   插图 {builder.fig_no} 张 ｜ 表格题注 {builder.tbl_no} 条 ｜ 目录条目 {len(builder.headings)} 条")
    print("   ✅ 文档文本中无 URL（https?://）" if not stats["urls"]
          else f"   ❌ 发现 URL：{stats['urls'][:5]}")
    if builder.missing_images:
        print(f"   ⚠️ 缺图：{builder.missing_images}")

    # ---- PDF ----
    outputs = [out_path]
    if not args.no_pdf:
        pdf = soffice_pdf(out_path, out_path.parent, work / "lo_profile")
        if pdf:
            print(f"✅ 生成 {pdf.relative_to(REPO)} ｜ {pdf_pages(pdf)} 页 ｜ "
                  f"{pdf.stat().st_size / 1024:.1f} KB")
            outputs.append(pdf)
        else:
            print("   ⚠️ PDF 转换失败（可手动用 Word / LibreOffice 导出）")
    _stamp_outputs(md_path, outputs)
    return 0


def _stamp_outputs(source: Path, outputs: list[Path]) -> None:
    """记录「产物 ← 源」的内容哈希（`docs/.generated.json`），供文档一致性测试校验。

    源改了却忘记重出 PDF/DOCX 是真实发生过的事故：旧产物会被一起提交，
    而 PDF 恰恰是评审最先读的形态。
    """
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from doc_stamp import write_stamp  # noqa: PLC0415

    write_stamp(source.parent / ".generated.json", source, [p for p in outputs if p])


def _blocks(md_text: str) -> list[Block]:
    """解析 + 表格题注 + 插图（每次都重新构造，避免两遍之间状态污染）。"""
    for fig in FIGURE_MAP:
        fig.used = False
    blocks = parse_markdown(md_text)
    attach_table_captions(blocks)
    return insert_figures(blocks, FIGURE_MAP)


if __name__ == "__main__":
    raise SystemExit(main())
