#!/usr/bin/env python
"""把 `docs/技术文档.md` 渲染成排版良好的 Word 与 PDF。

    .venv/bin/python scripts/build_tech_doc_pdf.py            # → docs/技术文档.docx + .pdf
    .venv/bin/python scripts/build_tech_doc_pdf.py --no-pdf   # 只出 docx

与 `build_report_docx.py` 同一套路（python-docx 排版 + LibreOffice 转 PDF），差异在于本脚本
面向**一般技术文档**：处理标题层级、段落、列表、表格、围栏代码块、插图与图注，并按章分页。
依赖：`python-docx`、`Pillow`（均为开发工具依赖）；`soffice` 缺失时只出 docx。
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
import sys
from pathlib import Path
from typing import Any, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = PROJECT_ROOT.parent
DEFAULT_SOURCE = REPO_ROOT / "docs" / "技术文档.md"

BODY_FONT = "Noto Sans CJK SC"
MONO_FONT = "DejaVu Sans Mono"
BODY_SIZE = 10.5
IMAGE_WIDTH_CM = 16.0
#: WebP 转 PNG 的缓存目录（不影响仓库内容）
CACHE_DIR = Path("/tmp/tech-doc-pdf-cache")
CAPTION_COLOR = RGBColor(0x55, 0x5F, 0x6B)


# --------------------------------------------------------------------------- #
# 解析：本脚本只需要 Markdown 的一个子集（标题 / 段落 / 列表 / 表格 / 代码块 / 插图）
# --------------------------------------------------------------------------- #
def parse_markdown(text: str) -> List[Tuple[str, Any]]:
    """把 Markdown 解析成 (块类型, 内容) 序列。

    支持：`#`~`####` 标题、`-`/`1.` 列表项、`|` 表格、``` 围栏代码、`![alt](src)` 插图、
    `<details>` 包裹块（跳过，避免把 Mermaid 源码印进正文）、其余为段落。
    """
    blocks: List[Tuple[str, Any]] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("<details>"):
            while i < len(lines) and "</details>" not in lines[i]:
                i += 1
            i += 1
            continue
        if not stripped or re.fullmatch(r"([-*_])\1{2,}", stripped):
            i += 1                                  # 空行与 Markdown 分隔线（---）跳过
            continue
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            body: List[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", (language, body)))
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if image:
            blocks.append(("image", (image.group(1), image.group(2))))
            i += 1
            continue
        if stripped.startswith("|"):
            table: List[List[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    table.append(cells)
                i += 1
            blocks.append(("table", table))
            continue
        if re.match(r"^([-*]|\d+\.)\s+", stripped):
            items: List[str] = []
            while i < len(lines) and re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]):
                items.append(re.sub(r"^\s*([-*]|\d+\.)\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("list", items))
            continue
        paragraph: List[str] = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.fullmatch(
                r"([-*_])\1{2,}", lines[i].strip()) and not re.match(
                r"^(#|\||```|!\[|[-*]\s|\d+\.\s|<details>)", lines[i].strip()):
            paragraph.append(lines[i].strip())
            i += 1
        blocks.append(("para", " ".join(paragraph)))
    return blocks


_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _materialize_image(path: Path, cache_dir: Path) -> Optional[Path]:
    """返回一张 python-docx 能嵌入的图片路径。

    界面截图是 WebP（体积小、适合仓库），而 python-docx 只认 PNG/JPEG 等常见格式，
    因此这里按需用 Pillow 转成 PNG 并缓存在临时目录，不改动仓库里的原图。
    """
    if path.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff"):
        return path
    try:
        from PIL import Image

        cache_dir.mkdir(parents=True, exist_ok=True)
        target = cache_dir / (path.stem + ".png")
        if not target.is_file() or target.stat().st_mtime < path.stat().st_mtime:
            with Image.open(path) as img:
                img.convert("RGB").save(target, "PNG")
        return target
    except Exception as exc:  # noqa: BLE001 - 转换失败时如实提示并跳过该图
        print(f"  [warn] 图片转换失败（{path.name}）：{exc}")
        return None


def _add_inline(paragraph: Any, text: str, *, size: float = BODY_SIZE,
                color: Optional[RGBColor] = None) -> None:
    """写入行内文本：粗体、行内代码与链接都按可读方式呈现。"""
    for token in _INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(size - 1)
        elif token.startswith("[") and "](" in token:
            label, target = token[1:-1].split("](", 1)
            run = paragraph.add_run(label if target.startswith("http") else f"{label}")
            if target.startswith("http"):
                run.font.color.rgb = RGBColor(0x1A, 0x5C, 0x8A)
        else:
            run = paragraph.add_run(token)
        if run.font.size is None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def _shade(paragraph: Any, fill: str = "F2F4F7") -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _page_number_footer(section: Any, label: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{label}    第 ")
    run.font.size = Pt(9)
    run.font.color.rgb = CAPTION_COLOR
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    tail = paragraph.add_run(" 页")
    tail.font.size = Pt(9)
    tail.font.color.rgb = CAPTION_COLOR


def build_docx(source: Path, out_docx: Path, *, title: str, subtitle: str) -> Path:
    blocks = parse_markdown(source.read_text(encoding="utf-8"))
    document = Document()

    style = document.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(BODY_SIZE)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style.paragraph_format.line_spacing = 1.4
    style.paragraph_format.space_after = Pt(6)

    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)
    _page_number_footer(section, title)

    # ---- 封面 ----
    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(150)
    run = cover.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_inline(sub, subtitle, size=12, color=CAPTION_COLOR)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_inline(meta, f"仓库：{REPO_URL}", size=10.5, color=CAPTION_COLOR)
    document.add_page_break()

    first_heading = True
    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload
            if level == 1 and not first_heading:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_heading = False
            heading = document.add_heading(level=min(level, 3))
            heading.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            heading.paragraph_format.space_after = Pt(6)
            run = heading.add_run(text)
            run.font.name = BODY_FONT
            run.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            run.font.size = Pt({1: 20, 2: 15, 3: 12.5}.get(level, 11))
            run.font.color.rgb = RGBColor(0x1F, 0x27, 0x33)
            continue
        if kind == "para":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            _add_inline(paragraph, payload)
            continue
        if kind == "list":
            for item in payload:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(2)
                _add_inline(paragraph, item)
            continue
        if kind == "code":
            language, body = payload
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            _shade(paragraph)
            for index, code_line in enumerate(body):
                run = paragraph.add_run(code_line)
                run.font.name = MONO_FONT
                run.font.size = Pt(9)
                if index < len(body) - 1:
                    run.add_break()
            if language:
                caption = document.add_paragraph()
                _add_inline(caption, f"代码：{language}", size=8.5, color=CAPTION_COLOR)
            continue
        if kind == "image":
            alt, src = payload
            raw_path = (source.parent / src).resolve()
            path = _materialize_image(raw_path, CACHE_DIR) if raw_path.is_file() else None
            if path is not None and path.is_file():
                picture = document.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.add_run().add_picture(str(path), width=Cm(IMAGE_WIDTH_CM))
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_inline(caption, f"图 {alt}", size=9, color=CAPTION_COLOR)
            else:
                print(f"  [warn] 插图缺失：{src}")
            continue
        if kind == "table":
            rows: Sequence[Sequence[str]] = payload
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, cell_text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    _add_inline(paragraph, cell_text, size=9.5)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
            document.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_docx))
    return out_docx


def soffice_pdf(docx: Path, out_dir: Path) -> Optional[Path]:
    """用 LibreOffice 转 PDF（与 build_report_docx.py 相同做法）。"""
    soffice = shutil.which("soffice") or "/usr/bin/soffice"
    if not Path(soffice).exists():
        print("  [warn] 未找到 soffice，跳过 PDF 转换")
        return None
    # 用户配置放系统临时目录：放到仓库里会污染工作区（首次生成时真的发生过）
    profile = Path(tempfile.gettempdir()) / "tech-doc-lo-profile"
    profile.mkdir(parents=True, exist_ok=True)
    cmd = [soffice, "--headless", "--norestore", "--invisible", "--nolockcheck",
           f"-env:UserInstallation=file://{profile}",
           "--convert-to", "pdf", "--outdir", str(out_dir), str(docx)]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=300)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        print(f"  [warn] soffice 转换失败：{exc}")
        return None
    pdf = out_dir / (docx.stem + ".pdf")
    return pdf if pdf.is_file() else None


def verify(pdf: Path) -> None:
    """打印页数与首页文本前若干行，确认中文与版面正常。"""
    size_kb = pdf.stat().st_size // 1024
    pages = "?"
    info = shutil.which("pdfinfo")
    if info:
        try:
            out = subprocess.run([info, str(pdf)], capture_output=True, text=True, timeout=60).stdout
            match = re.search(r"^Pages:\s+(\d+)", out, re.M)
            if match:
                pages = match.group(1)
        except (OSError, subprocess.TimeoutExpired) as exc:  # 允许静默：页数只是提示信息
            print(f"  [warn] pdfinfo 读取页数失败：{exc}")
    head = ""
    text_tool = shutil.which("pdftotext")
    if text_tool:
        try:
            out = subprocess.run([text_tool, "-f", "1", "-l", "2", str(pdf), "-"],
                                 capture_output=True, text=True, timeout=60).stdout
            head = " ".join(out.split())[:120]
        except (OSError, subprocess.TimeoutExpired) as exc:  # 允许静默：仅用于打印自检
            print(f"  [warn] pdftotext 自检失败：{exc}")
    print(f"  PDF：{pdf}（{size_kb} KB，{pages} 页）")
    if head:
        print(f"  首页文本：{head}")


REPO_URL = "https://github.com/yuhao233/Docking-Multi-agent"


def _stamp_outputs(source: Path, outputs: List[Path]) -> None:
    """记录「产物 ← 源」的内容哈希（`docs/.generated.json`）。

    供 `tests/test_docs_consistency.py` 校验产物是否与源同步 —— 源改了却忘记重出
    PDF/DOCX 是真实发生过的事故（旧产物会被一起提交）。
    """
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from doc_stamp import write_stamp  # noqa: PLC0415

    write_stamp(source.parent / ".generated.json", source, [p for p in outputs if p])


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="把技术文档渲染为 Word 与 PDF")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE))
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args(argv)

    source = Path(args.source)
    if not source.is_file():
        print(f"找不到源文件：{source}", file=sys.stderr)
        return 2
    title = "分子对接多 Agent 协作系统"
    subtitle = "技术文档 · 框架 / 运行原理 / 使用说明"
    docx_path = source.with_suffix(".docx")
    print(f"[tech-doc] 解析 {source}")
    build_docx(source, docx_path, title=title, subtitle=subtitle)
    print(f"  已生成 {docx_path}（{docx_path.stat().st_size // 1024} KB）")
    if args.no_pdf:
        _stamp_outputs(source, [docx_path])
        return 0
    pdf = soffice_pdf(docx_path, source.parent)
    if pdf is None:
        print("  未生成 PDF（只保留 docx）")
        _stamp_outputs(source, [docx_path])
        return 1
    verify(pdf)
    _stamp_outputs(source, [docx_path, pdf])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
